#!/usr/bin/env python3
"""C5h: score David's completed Wisconsin sheet into verdicts, joined to our internal
flags/signals so outcome can be read against fit_flag / alek_signal / lien age / lender.

David's WI grading (decoded 2026-09-14) lives on the *Comments* cell, not the row:
  green  FF92D050  -> positive / maybe (open door, callback with intent)
  blue   theme 7   -> followup (call back, timing/gatekeeper)
  light  theme 5   -> existing / CRM (paired with LP/GM/Last Contact filled; "DO NOT CALL")
  yellow FFFFFF00  -> dead / not interested
  (red font on Comments = harder "not interested"; on Name/Phone = data-quality flag)
  grey   theme 2   -> "info not provided, I added it" (data marker, NOT an outcome)
LP/GM columns: his CRM-presence flags, only "Yes" on the 5 existing rows.

Inputs  data/raw/david_wi_completed_2026-09-14.xlsx  (David, called all 50)
        out/WI_PHONE_TEST_50_2026-09-11.csv          (our internal, flags+signals)
Output  out/verdicts/david_wi_2026-09-14.csv
"""
import csv, re
import openpyxl

XLSX = "data/raw/david_wi_completed_2026-09-14.xlsx"
INT = "out/WI_PHONE_TEST_50_2026-09-11.csv"
OUT = "out/verdicts/david_wi_2026-09-14.csv"

# comment-cell fill -> david_colour
GREEN, YELLOW = "FF92D050", "FFFFFF00"


def norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def fill_code(cell):
    """Return one of: green, yellow, blue, existing, grey, '' for a cell's fill."""
    f = cell.fill
    if not f.patternType:
        return ""
    fg = f.fgColor
    if fg.type == "rgb":
        if fg.rgb == GREEN:
            return "green"
        if fg.rgb == YELLOW:
            return "yellow"
        return ""
    if fg.type == "theme":
        return {7: "blue", 5: "existing", 2: "grey", 4: "grey"}.get(fg.theme, "")
    return ""


def font_red(cell):
    fc = cell.font.color
    return bool(fc and getattr(fc, "type", None) == "rgb" and getattr(fc, "rgb", None) == "FFFF0000")


# heuristic contact detection from David's comment shorthand
NO_CONTACT = re.compile(r"\bLVM\b|\bLMIG\b|\bLMWS\b|\bLMIGMB\b|\bLM\b|voicemail|no answer|ringer|straight to mb|call is blocked|nis\b|vm is not", re.I)
REACHED = re.compile(r"\bsaid\b|hanged up|hung up|screened|answered|short|transfer|didn.t want|not interested|vacation|no current needs|looking at|zero interest|wouldn.t", re.I)
OWNER_WORDS = re.compile(r"\bhe\b|\bhim\b|\bhis\b|owner|said no|looking at|zero interest|screened his own", re.I)


def outcome_of(colour, red, comment):
    c = comment.lower()
    if colour == "existing" or "do not call" in c:
        return "existing_lead"
    if colour == "green":
        return "maybe_later" if ("call back" in c or "callback" in c or "ok with" in c or "next" in c) else "followup"
    if colour == "blue":
        return "followup"
    if colour == "yellow" or red:
        return "dead"
    # no colour: LVM / offered promo / could not reach
    return "no_decision"


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[wb.sheetnames[0]]
    hdr = [c.value for c in ws[1]]
    ix = {h: i + 1 for i, h in enumerate(hdr) if h}

    dav = {}  # norm(company) -> dict
    for r in range(2, ws.max_row + 1):
        comp = ws.cell(row=r, column=ix["Company"]).value
        if not comp or str(comp).startswith("Info was not"):
            continue
        cm = ws.cell(row=r, column=ix["Comments"])
        comment = (cm.value or "").strip()
        colour = fill_code(cm)
        # existing cluster: LP/GM light-filled even if comment cell isn't
        lp_fill = fill_code(ws.cell(row=r, column=ix["LP"]))
        if lp_fill == "existing":
            colour = "existing"
        red = font_red(cm)
        name_added = fill_code(ws.cell(row=r, column=ix["Name"])) == "grey"
        outcome = outcome_of(colour, red, comment)
        reached = bool(REACHED.search(comment)) and not NO_CONTACT.match(comment.strip())
        # "he/his said ... / looking at / zero interest" => owner-level
        owner = reached and bool(OWNER_WORDS.search(comment)) and outcome != "existing_lead"
        dav[norm(comp)] = {
            "lp": (ws.cell(row=r, column=ix["LP"]).value or "").strip(),
            "gm": (ws.cell(row=r, column=ix["GM"]).value or "").strip(),
            "last_contact": (str(ws.cell(row=r, column=ix["Last Contact"]).value or "")[:10]),
            "david_colour": colour or ("red_font" if red else ""),
            "outcome": outcome, "human_reached": reached, "owner_reached": owner,
            "name_added": name_added, "david_comment": comment,
        }

    internal = list(csv.DictReader(open(INT)))
    out = []
    miss = []
    for r in internal:
        d = dav.get(norm(r["Company"]))
        if not d:
            miss.append(r["Company"]); continue
        crm = "yes" if (d["lp"].lower().startswith("y") or d["gm"].lower().startswith("y")) else "no"
        out.append({
            "company": r["Company"], "lane": r["lane"], "line_type": r["line_type"],
            "cnam_match": r["cnam_match"], "fit_flag": r["fit_flag"], "alek_signal": r["alek_signal"],
            "lien_age_months": r["lien_age_months"], "ucc_status": r["ucc_status"],
            "lender": r["ucc_secured_party"], "lender_class": r["ucc_lender_class"],
            "power_units": r["power_units"], "crm": crm,
            "lp": d["lp"], "gm": d["gm"], "last_contact": d["last_contact"],
            "david_colour": d["david_colour"], "outcome": d["outcome"],
            "human_reached": d["human_reached"], "owner_reached": d["owner_reached"],
            "name_added": d["name_added"], "david_comment": d["david_comment"],
        })
    with open(OUT, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=out[0].keys()); w.writeheader(); w.writerows(out)

    from collections import Counter
    print(OUT, len(out), "rows;  unmatched:", miss)
    print("outcome:", dict(Counter(o["outcome"] for o in out)))
    print("david_colour:", dict(Counter(o["david_colour"] for o in out)))
    print("crm:", dict(Counter(o["crm"] for o in out)))
    print("human_reached:", sum(o["human_reached"] for o in out), " owner_reached:", sum(o["owner_reached"] for o in out))
    print("name_added(no owner sent):", sum(o["name_added"] for o in out))


if __name__ == "__main__":
    main()
