#!/usr/bin/env python3
"""C5j: score David's *partial* FL-100 and CO-100 sheets into verdicts, joined to the
internal files that still carry intent (alek_signal / fit_flag / lien age / lender).

David returned both before finishing (call notes on the call before ours):
  FL  rows 2-51  = first 50 of 100 called   (blind copy -> intent columns were stripped
                   from his sheet; we join back to the internal FL_PHONE_TEST_100 file)
  CO  rows 2-29  = first 28 of 99 called
  flc = the completed FL 100 (2026-09-19); theme8 = funded in that workbook

Colour scheme is the WI scheme but each workbook uses a different theme palette, so we map
by *canonical meaning*, confirmed against the comment text:
  green   FF92D050 / FL theme6      -> positive / callback-with-intent   (maybe_later)
  blue    WI theme7 / CO theme3     -> followup / timing / try-again      (followup)
  yellow  FFFFFF00                  -> dead / not interested
  red     FFFF0000 fill or font     -> hard dead ("take me off the list")
  existing WI/theme5                -> CRM lead ("DO NOT CALL, X's lead")  (existing_lead)
  funded  CO theme9                 -> CRM funded customer                 (existing_lead, funded=1)
  grey    theme2/4                  -> data marker, not an outcome
Only rows David actually annotated are scored; blank rows below his stop point are skipped.

Usage:  .venv/bin/python scripts/c5j_verdicts.py fl
        .venv/bin/python scripts/c5j_verdicts.py co
"""
import csv, re, sys
import openpyxl

CFG = {
    "fl": dict(
        xlsx="data/raw/david_fl_partial_2026-09-15.xlsx",
        internal="out/FL_PHONE_TEST_100_2026-09-14.csv",
        out="out/verdicts/david_fl_partial_2026-09-15.csv",
        comments_col="Comments",
    ),
    "flc": dict(
        xlsx="data/raw/david_fl_completed_2026-09-19.xlsx",
        internal="out/FL_PHONE_TEST_100_2026-09-14.csv",
        out="out/verdicts/david_fl_completed_2026-09-19.csv",
        comments_col="Comments",
    ),
    "co": dict(
        xlsx="data/raw/david_co_partial_2026-09-15.xlsx",
        internal="out/CO_PHONE_TEST_100_2026-09-15.csv",
        out="out/verdicts/david_co_partial_2026-09-15.csv",
        comments_col="Comments",
    ),
}

GREEN, YELLOW, RED = "FF92D050", "FFFFFF00", "FFFF0000"


def norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def fill_code(cell):
    """Canonical fill label across both workbooks' theme palettes."""
    f = cell.fill
    if not f.patternType:
        return ""
    fg = f.fgColor
    if fg.type == "rgb":
        return {GREEN: "green", YELLOW: "yellow", RED: "red"}.get(fg.rgb, "")
    if fg.type == "theme":
        return {7: "blue", 3: "blue", 5: "existing", 9: "funded", 8: "funded",
                6: "green", 2: "grey", 4: "grey"}.get(fg.theme, "")
    return ""


def font_red(cell):
    fc = cell.font.color
    return bool(fc and getattr(fc, "type", None) == "rgb" and getattr(fc, "rgb", None) == RED)


# comment shorthand -> did a live human pick up?  (LM* / VM / MB / NIS / screened-by-AI = no)
NO_CONTACT = re.compile(
    r"\bL[MV]\w*\b|voicemail|straight to (?:a )?\w*\s*mb|generic mb|mb is full|"
    r"\bnis\b|no answer|ai assi|ai assistant|blocked|spanish mb|dropped the call",
    re.I)
REACHED = re.compile(
    r"\bsaid\b|hanged up|hung up|picked up|answered|screened (?:his|the|her)|"
    r"not interested|not looking|not buying|all set|no plans|no need|selling|getting out|"
    r"take[n]? (?:me )?off|didn.?t want|full fleet|we need nothing|set now|out till|"
    r"back in|call back|nlwc|is not in", re.I)
OWNER = re.compile(
    r"\bhe (?:said|goes|is)\b|\bhis\b|owner|screened his own|"
    r"said (?:no|they|we)|all set|no plans|getting out|set now|nlwc", re.I)


def outcome_of(colour, red, comment):
    c = comment.lower()
    if colour in ("existing", "funded") or "do not call" in c:
        return "existing_lead"
    if colour == "green":
        return "maybe_later"
    if colour == "blue":
        return "followup"
    if colour in ("yellow", "red") or red:
        return "dead"
    # no colour: leaned message / could not reach / AI-screened
    return "no_decision"


def score_sheet(cfg):
    wb = openpyxl.load_workbook(cfg["xlsx"])
    ws = wb[wb.sheetnames[0]]
    hdr = [c.value for c in ws[1]]
    ix = {h: i + 1 for i, h in enumerate(hdr) if h}
    ccol = ix[cfg["comments_col"]]

    dav = {}
    for r in range(2, ws.max_row + 1):
        comp = ws.cell(row=r, column=ix["Company"]).value
        if not comp:
            continue
        cm = ws.cell(row=r, column=ccol)
        comment = (cm.value or "").strip()
        colour = fill_code(cm)
        # existing/funded may be flagged on LP/GM instead of the comment cell
        for k in ("LP", "GM"):
            if k in ix:
                fc = fill_code(ws.cell(row=r, column=ix[k]))
                if fc in ("existing", "funded"):
                    colour = colour or fc
        red = font_red(cm)
        if not (comment or colour or red):
            continue  # David hasn't reached this row yet
        outcome = outcome_of(colour, red, comment)
        reached = bool(REACHED.search(comment)) and not NO_CONTACT.match(comment.strip())
        owner = reached and bool(OWNER.search(comment)) and outcome != "existing_lead"
        dav[norm(comp)] = {
            "lp": str(ws.cell(row=r, column=ix["LP"]).value or "").strip() if "LP" in ix else "",
            "gm": str(ws.cell(row=r, column=ix["GM"]).value or "").strip() if "GM" in ix else "",
            "david_colour": colour or ("red_font" if red else ""),
            "funded": 1 if colour == "funded" else 0,
            "outcome": outcome, "human_reached": int(reached), "owner_reached": int(owner),
            "david_comment": comment,
        }
    return dav


def main():
    which = sys.argv[1].lower()
    cfg = CFG[which]
    dav = score_sheet(cfg)

    internal = list(csv.DictReader(open(cfg["internal"])))
    keep = ["fit_flag", "alek_signal", "lender_note", "lane", "line_type", "cnam_match",
            "lien_age_months", "ucc_status", "ucc_secured_party", "ucc_lender_class",
            "power_units", "equipment_ticket"]
    out, miss, seen = [], [], set()
    for r in internal:
        key = norm(r["Company"])
        d = dav.get(key)
        if not d or key in seen:
            continue  # row David hasn't called yet, or a duplicate internal row
        seen.add(key)
        row = {"company": r["Company"]}
        for k in keep:
            row[k] = r.get(k, "")
        crm = "yes" if (d["lp"].lower().startswith("y") or d["gm"].lower().startswith("y")) else "no"
        row.update({"crm": crm, "funded": d["funded"], "david_colour": d["david_colour"],
                    "outcome": d["outcome"], "human_reached": d["human_reached"],
                    "owner_reached": d["owner_reached"], "david_comment": d["david_comment"]})
        out.append(row)
    matched = {norm(r["Company"]) for r in internal}
    miss = [dav[k] for k in dav if k not in matched]

    with open(cfg["out"], "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=out[0].keys()); w.writeheader(); w.writerows(out)

    from collections import Counter
    print(cfg["out"], len(out), "scored;  unmatched-in-internal:", len(miss))
    print("outcome:", dict(Counter(o["outcome"] for o in out)))
    print("david_colour:", dict(Counter(o["david_colour"] for o in out)))
    print("human_reached:", sum(o["human_reached"] for o in out),
          " owner_reached:", sum(o["owner_reached"] for o in out),
          " funded:", sum(o["funded"] for o in out), " crm=yes:", sum(o["crm"] == "yes" for o in out))


if __name__ == "__main__":
    main()
